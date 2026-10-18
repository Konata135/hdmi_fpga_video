"""Dump a .docx's paragraphs and tables to UTF-8 text (the console here is GBK).
Usage: python tools/docx_to_text.py <in.docx> <out.txt>
Kept next to pdf_to_text.py: the competition and vendor documents are .docx and
.pdf, and both are easier to search once they are text.
"""
import sys
from docx import Document


def main():
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(src)
    with open(dst, "w", encoding="utf-8") as f:
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                f.write(text + "\n")
        for ti, table in enumerate(doc.tables):
            f.write(f"\n===== table {ti + 1} =====\n")
            for row in table.rows:
                cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                f.write(" | ".join(cells) + "\n")
    print("written", dst)


if __name__ == "__main__":
    main()
